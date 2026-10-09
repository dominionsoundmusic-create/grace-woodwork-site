#!/usr/bin/env python3
"""Keyword plan for gracewoodworkkilgore.com.

Reads keywords/Grace-Woodwork-Keywords.xlsx (ALL tab), keeps every keyword with 10+ searches,
re-sorts column H ("Fit for Grace") with the rules below, and gives every keyword that fits Grace
either a page (target or supporting) or a one-line reason it is deliberately unused.

Rules are tried in order; the first match wins. Every keyword on the "FIT FOR GRACE 10+" tab must
be matched by an explicit rule (not a fallback), so nothing that fits is assigned by accident.

Writes docs/keyword-plan.md and docs/keyword-assignments.csv.
Usage: python3 scripts/keyword_plan.py
"""
import csv
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "keywords" / "Grace-Woodwork-Keywords.xlsx"

PAGES = {
    "home": ("/", "Homepage"),
    "fr": ("/furniture-repair/", "Furniture repair (hub)"),
    "cf": ("/custom-furniture/", "Custom furniture (hub)"),
    "cc": ("/custom-cabinets/", "Custom cabinets (hub)"),
    "cs": ("/custom-signs/", "Custom signs (hub)"),
    "paint": ("/custom-cabinets/cabinet-painting/", "Cabinet painting"),
    "reface": ("/custom-cabinets/cabinet-refacing/", "Cabinet refacing"),
    "builtin": ("/custom-cabinets/built-in-shelves/", "Built-in shelves and bookshelves"),
    "refinish": ("/furniture-repair/furniture-refinishing/", "Furniture refinishing"),
    "antique": ("/furniture-repair/antique-furniture-restoration/", "Antique furniture restoration"),
    "chair": ("/furniture-repair/chair-caning-and-repair/", "Chair caning and chair repair"),
    "table": ("/furniture-repair/table-refinishing/", "Table refinishing"),
    "dining": ("/custom-furniture/custom-dining-tables/", "Custom dining tables"),
    "deck": ("/deck-repair/", "Deck repair"),
    "porch": ("/porch-repair/", "Porch repair"),
    "g-paintcost": ("/custom-cabinets/cabinet-painting-cost/", "Guide: what cabinet painting costs"),
    "g-reface": ("/custom-cabinets/what-is-cabinet-refacing/", "Guide: what cabinet refacing is"),
    "g-refcost": ("/furniture-repair/furniture-refinishing-cost/", "Guide: what furniture refinishing costs"),
    "g-veneer": ("/furniture-repair/furniture-veneer-repair/", "Guide: fixing furniture veneer"),
    "g-cane": ("/furniture-repair/how-chair-caning-is-repaired/", "Guide: how chair caning is repaired"),
    "g-table8": ("/custom-furniture/dining-table-size-for-8/", "Guide: what size dining table seats 8"),
    "gallery": ("/gallery/", "Gallery"),
}
for svc, hub in (("furniture-repair", "fr"), ("custom-furniture", "cf"), ("custom-cabinets", "cc"), ("custom-signs", "cs")):
    for town in ("kilgore", "longview", "tyler", "henderson", "marshall", "gladewater", "white-oak", "overton", "troup",
                 "lindale", "big-sandy", "gilmer", "arp", "new-london", "hallsville"):
        PAGES[f"{svc}:{town}"] = (f"/{svc}/{town}-tx.html", f"{svc.replace('-', ' ').capitalize()} in {town.replace('-', ' ').title()} (town page, text unchanged)")

TARGETS = {
    "home": "woodworking shop", "fr": "furniture repair", "cf": "custom furniture", "cc": "custom cabinets",
    "cs": "custom wood signs", "paint": "cabinet painting", "reface": "cabinet refacing", "builtin": "built in shelves",
    "refinish": "furniture refinishing", "antique": "antique furniture restoration", "chair": "chair caning",
    "table": "table refinishing", "dining": "custom dining table", "deck": "deck repair", "porch": "porch repair",
    "g-paintcost": "cabinet painting prices", "g-reface": "what is cabinet resurfacing",
    "g-refcost": "furniture refinishing cost", "g-veneer": "how to repair furniture veneer",
    "g-cane": "how to repair chair caning", "g-table8": "what size dining table seats 8",
}

# Places that are not Grace's East Texas service area. A keyword naming one is "Other place".
ELSEWHERE = r"""utah|colorado|denver|new jersey|tallahassee|arizona|phoenix|north carolina|vermont|green bay|alpharetta|
naperville|charlottesville|eugene|oregon|ocala|rhode island|fairfax|fort collins|kalamazoo|kalispell|myrtle beach|salem|
\bbend\b|elk grove|hawaii|oahu|maine|michigan|minnesota|new hampshire|ohio|palm desert|puyallup|traverse city|cape cod|
santa barbara|vero beach|daytona|evanston|roseville|\byork\b|maryland|bethesda|escondido|kenosha|ventura|santa cruz|
torrance|walnut creek|yakima|bozeman|california|swainsboro|georgia|virginia|edinburgh|barrie|jupiter|mississauga|
victoria|winnipeg|etobicoke|gold coast|halifax|hobart|huntington beach|kelowna|kingston|quad cities|massachusetts|
abu dhabi|indonesia|\bnz\b|oakville|jaipur|pennsylvania|burlington|downers grove|hamilton|twin cities|vaughan|
brantford|fredericton|kamloops|lethbridge|newmarket|regina|thunder bay|windsor|culver city|encinitas|qatar|
west chester|youngstown|johannesburg|\bkent\b|kildare|liverpool|newcastle|oxford|hastings|yorkshire|zeewolde|
the villages|china|ghana|philippines|washington|tacoma|connecticut|el paso|midland|san antonio|austin|houston|
los angeles|dallas|fort worth|chicago|atlanta|seattle|portland|miami|orlando|tampa|nashville|charlotte|raleigh|
boston|las vegas|san diego|sacramento|kansas|missouri|st louis|indianapolis|columbus|cleveland|cincinnati|detroit|
pittsburgh|philadelphia|baltimore|richmond|jacksonville|florida|louisiana|oklahoma|arkansas|tennessee|alabama|
mississippi|kentucky|indiana|illinois|wisconsin|iowa|nebraska|idaho|montana|wyoming|nevada|new mexico|new york|
\bnyc\b|brooklyn|long island|toronto|ontario|calgary|edmonton|vancouver|ottawa|montreal|canada|\buk\b|london|
australia|sydney|melbourne|brisbane|perth|adelaide|dubai|india|singapore|ireland|scotland|dublin|hendersonville|
\bnc\b|\bfl\b|\baz\b|\bnj\b|\bny\b|\bid 83423|kilgore id|bay area|silicon|honolulu|anchorage|alaska|spokane|boise|tucson|mesa|scottsdale|
albuquerque|el monte|jamaica|riverside|irvine|san jose|oakland|fresno|bakersfield|reno|salt lake|provo|ogden|
lubbock|amarillo|corpus christi|waco|killeen|abilene|odessa|beaumont|mcallen|brownsville|laredo|plano|frisco|
arlington|irving|garland|mckinney|denton|round rock|katy|sugar land|the woodlands|conroe|pearland|league city|
galveston|college station|bryan|temple|san marcos|new braunfels|georgetown|cedar park|pflugerville|lufkin|
nacogdoches|texarkana|sherman|paris tx|greenville|athens tx|palestine|jacksonville tx|mount pleasant|sulphur springs|
canton tx|mineola|quitman|winnsboro|carthage|center tx|jefferson tx|daingerfield|pittsburg tx|atlanta tx""".replace("\n", "")
EAST_TX = r"\b(kilgore(?! id)|longview|tyler|henderson(?!ville)|marshall|gladewater|white oak|overton|troup|lindale|big sandy|gilmer|\barp\b|new london|hallsville|east texas|gregg county|smith county|rusk county)\b"

# (pattern, fit, page key or None, role or reason)
# role: "target" / "supporting". When page is None, the last field is the reason it is unused.
R = []


def rule(pat, fit, page, note=""):
    R.append((re.compile(pat), fit, page, note))


SVC, NEAR, Q, ET, ELSE, DIY, OFF, BRAND = ("Service", "Service near me", "Question", "East Texas", "Other place",
                                          "DIY or product", "Off-topic", "Other business or brand")

# ---- 1. not Grace: other businesses, brands, shops and noise ------------------------------------------------------
rule(r"sherwin|benjamin moore|\bbehr\b|klingspor|home depot|lowe'?s|ikea|cabinets to go|cabinets\.com|restoration hardware|"
     r"ethan allen|\btarget\b|ebay|bunnings|disney|quatrine|thomas johnson|tom johnson|zimmerman|zongkers|wolfe'?s|"
     r"us cabinet refacing|quality cabinet refacing|antiques & furniture restoration inc|antique furniture repair llc|"
     r"fs antique|easy does it|chair repair shoppe|furniture repair 911|sofa repair zone|wood wright|yankee woodworking|"
     r"custom furniture expressions|custom cabinets by lawrence|chair caning by anne|picasso|stardew|abiotic|"
     r"carpentry 7th|carpentry xp|woodshop 510|^custom 9$|^chair \d|^furniture \d|^8 furniture|chair 1 barber|"
     r"kitchen cabinets kings|kitchen cabinets express|kitchen cabinets direct|kitchen cabinets 84 lumber|cabinet xr|"
     r"cabinet x paint|ace hardware|rejuvenate|furniture repair bank|table rock restoration|trendy wood shop|"
     r"woodworking shop com|woodworking shop inc|custom cabinets inc$|custom cabinets llc|custom cabinets com|"
     r"furniture repair yelp|hsn code|facebook marketplace|kitchen cabinets near me open now|chair caning picasso|"
     r"custom furniture solutions|custom cabinets by design|furniture restoration inc|quality cabinet refacing inc|"
     r"\bmenards\b|wayfair|amazon|costco|walmart|lumber liquidators|cabinet depot|furniture row|rooms to go|ashley",
     BRAND, None, "Another company's name or a brand: searchers want that business, not Grace.")
rule(r"3/8 zodiac|table painting (?:designs|art)|table painting ideas|street sign color|construction signs|chairs explode|what chair can move|^where chair$|^when furniture$|"
     r"^are chair$|^how made furniture$|who made furniture|^made furniture quality$|how wood charity|^will wood shop$|"
     r"^custom wood sizes$|z shelves|table fan repair|direction to tyler|is tyler texas conservative|caleb kilby|cameron kilgore|"
     r"^kilgore tx 7566|kilgore tx 75662|signing time|carpenteria|^furniture tyler|tyler texas furniture stores|"
     r"home furnishings tyler|henderson furniture reviews|^furniture kilgore|^kilgore furniture$|custom homes tyler|"
     r"furniture tyler tx|how many restoration hardware",
     OFF, None, "Off-topic or noise: not a search for woodwork Grace does.")
rule(r"\bjobs?\b|salary|reddit|apprenticeship|pensacola|kitchener|lahore|bangalore|classes|\bclass\b|course|youtube|videos?\b|chair lift|stair ?lift|revolving chair|union|work pants|"
     r"in shower|for shower|next to shower|dog crate|for restaurants|price in sri lanka|cane chair used|for sale$|"
     r"will furniture stores|furniture restoration of marin|joanna gaines|how many supports for a deck",
     OFF, None, "Jobs, classes, videos, stairlifts or other searches that are not a customer hiring a woodwork shop.")
rule(r"free estimate", SVC, None, "Grace has not confirmed free estimates, and the site does not use the phrase.")
rule(r"furniture repair (?:in home|mobile|at home|home service|on site)|chair repair (?:at home|home service)|"
     r"furniture repair near me in home", SVC, None, "In-home repair: Grace works in its Kilgore shop and has not confirmed on-site furniture repair.")
rule(r"\bwax\b|sticks\b|repair hardware|furniture hardware|deck repair filler", DIY, None, "A product, tool or material for doing it yourself: product pages win these, not a shop.")
# ---- 2. services Grace does not offer -------------------------------------------------------------------------------
rule(r"upholster|leather|recliner|zipper|springs?\b|hydraulic|wheel ?chair|office chair|chair office|seat repair patch|"
     r"furniture repair patches|chair replacement fabric|cushion|slipcover|covers?\b|jack repair|webbing|straps|"
     r"furniture repair welding|chair repair plastic|seat repair kit|chair cover",
     OFF, None, "Upholstery, leather, mechanisms or covers: not a service Grace has confirmed.")
rule(r"floor", OFF, None, "Floor refinishing: not a service Grace has confirmed.")
rule(r"screen|brick|concrete|foundation|porch replacement windows|porch lift|porch roof|porch ceiling|umbrella|"
     r"table saw top|deck rebuild|replace deck|how much replace deck|porch replacement",
     OFF, None, "Screens, masonry, roofing or full replacement: outside the porch and deck repair Grace confirmed.")
# ---- 3. hobby shops, tools, DIY products and how-to that would need unsafe or unhelpful content ----------------------
rule(r"dust collection|vac dust|vacuum system|air filt|air clean|air purif|jigs|for rent|rental|space for rent|membership|"
     r"woodworking shop (?:for beginners|projects|table|organization|designs|setup|equipment|storage|build|heater|pictures|"
     r"cart|essentials|flooring|lights|layout|size|tours|auction|garage|must haves|names|planner|hacks|rate|the villages|"
     r"victoria|coat|floor plans|tools|stool|table plans|storage ideas|furniture|cabinets|signs|for sale)|wood shop (?:apron|organization|work bench|"
     r"work table|online|utah|gifts|work$|ideas)|woodworkers shop|12x20|what is a good size for a woodworking|"
     r"set up a small woodworking|woodworking examples|is woodworking|learn woodworking|what woodworking items|"
     r"where to sell|how to sell|sell custom|how much do custom furniture makers make|furniture restoration business|"
     r"furniture refinishing business|is furniture restoration profitable|custom furniture business|custom furniture website|"
     r"custom furniture logo|custom furniture tags|custom wood signs business|woodworking supply|woodworking shop near me open|"
     r"wood shop air|^what is wood shop$|carpenter class|furniture repair training|furniture restoration training|"
     r"furniture restoration (?:books|videos|workshop$)|furniture repair technician$",
     DIY, None, "Hobby woodworking shops, tools or the business of woodworking: not a customer looking for a shop to hire.")
rule(r"\bkit\b|markers?\b|crayons|pens\b|wax sticks|putty|paste|wood filler|glue|tape\b|supplies|sander|tools?\b|"
     r"products|polish|primer|\bspray\b|roller|brush|hangers|stands?\b|drying rack|painting rack|hooks|gallon|one coat|"
     r"no prep|no sanding|requires no|peel and stick|vinyl wrap|for laminate|on walls|be used on walls|looks like wood|"
     r"for furniture$|cabinet paint for bathroom|water based|with hardener|hardener|touch up|epoxy|parts\b|"
     r"chair cane plastic|chair caning book|where to buy chair caning|repair materials|furniture repair materials|"
     r"deck repair materials|antique furniture restoration materials|furniture restoration equipment|"
     r"furniture refinishing equipment|cabinet paint (?:at|gloss|that|on)|cabinet joint paint|furniture restoration oil|"
     r"furniture refinishing oil|how much cabinet paint do i need|where to buy cabinet paint|cabinet paint options|"
     r"is cabinet paint different|cabinet paint enamel|can cabinet paint|paint cabinets yourself",
     DIY, None, "A product, tool or material for doing it yourself: product pages win these, not a shop.")
rule(r"how to (?:build|make|install)|diy|\bplans\b|how to$|how to repair (?:deck|porch|furniture$|broken|chair upholstery|chair hydraulic)|"
     r"how to repair deck|how to repair porch|refinish table how to|cabinet painting how to|furniture refinishing how to|"
     r"antique furniture restoration how to|chair caning how to|deck repair how to|caning chairs for beginners|"
     r"for beginners|techniques|tips\b|ideas\b|decor|refacing cabinets yourself|how to make a wooden sign|"
     r"how to make custom wood signs|how restore furniture|how refurbish furniture|how to repair furniture scratch|"
     r"how to repair wood furniture finish|chair caning directions|between studs|drywall|what to put on",
     DIY, None, "Do-it-yourself how-to: the plan does not publish step-by-step DIY repair, power-tool or structural instructions; pages answer the question and say when to call.")
# ---- 4. shopping for new mass-market products ----------------------------------------------------------------------
rule(r"cane chair (?:dining|black|high back|living room|outdoor|cushions?|table set|2 seater|3d|for office|online|"
     r"hanging|lounge|disney|nz|bunnings|target|ebay)|cane chairs ghana|ebay cane chair|cane chair dining set|"
     r"antique (?:chair|furniture) for restoration|old furniture for restoration|custom furniture covers",
     DIY, None, "Shopping for a new or project piece, not for repair.")
rule(r"kitchen cabinets? (?:in stock|store|stores|for sale|on sale|sale|used|wholesale|outlet|outlets|liquidators|"
     r"cheap|for cheap|clearance|on clearance|discount|distributors|online|order online|rta|ready to assemble|unassembled|"
     r"assembled|full set|set$|10x10|kit|from china|brands|manufacturers|makers$|best price|on a budget|near$|nearby|"
     r"near me cheap|near me for sale|hardware|knobs|pulls|handles|hinges|locks|liners|mats|cleaner|covers|holders|"
     r"brackets|parts|names|quotation|estimate|usa|jamaica|el monte|philippines|2025|2026|trends|images|pictures|gallery|"
     r"examples|design tool|layout|dimensions|standard sizes|depth|\d+ ?(?:inch|inches|ft|foot|deep|wide)|\d+$|"
     r"for mobile homes|toe kick|crown molding|molding|trim|end panels|top$|bottom$|base$|upper$|wall$|tall$|high$|"
     r"for island|corner|with sink|with countertops|countertops|on legs|under sink|under lights|with drawers|"
     r"pull out|stand alone|for storage|storage|extra storage|for small|small|ideas|design|styles|types|materials|"
     r"quality$|modern|european|shaker|vintage|glass|glossy|metal|laminate|vinyl|unfinished|natural wood|light wood|"
     r"knotty pine|oak$|maple$|wood$|solid wood|white|black|green|blue|navy|gray|grey|beige|brown|yellow|espresso|"
     r"light gray|colors?|color|2 tone|two colors|2 different colors|paint colors|paint ideas|paint$|gold handles|"
     r"no handles|with handles|that go to the ceiling|all the way|vaulted|\d+ ?ft|9 foot|new$|outdoor|garage|"
     r"in bathroom|pantry|shelves|for sale cheap|installers|installed$|installation cost|replacement cost|replacement$|"
     r"remodel|that look like furniture|cost$|near me$)|^kitchen cabinets$|where (?:to )?buy kitchen|where to get kitchen|"
     r"where to buy kitchen storage|who sale kitchen|kitchen cabinets? knobs|which kitchen cabinets|what kitchen cabinets|"
     r"what kitchen cabinet colors|why white kitchen|will white kitchen|how many kitchen cabinets|how kitchen cabinets are|"
     r"are kitchen cabinets|where are kitchen cabinets|how much (?:kitchen cabinets|will kitchen cabinets)|"
     r"how much kitchen cabinets|when to replace kitchen cabinets|can kitchen cabinets be removed|who install kitchen|"
     r"does cabinets to go|are cabinets to go|is cabinets|semi custom cabinets near me",
     DIY, None, "Shopping for stock kitchen cabinets, hardware or design ideas: retailers and showrooms own these; Grace builds and refinishes, it does not sell stock lines.")
rule(r"custom cabinets? (?:cheap|online|rta|wholesale|diy|ikea|made to order$)|custom cabinets online|"
     r"custom furniture (?:stores|china|online|outlet|factory|brands|manufacturers|abu dhabi|indonesia|risers|"
     r"design online|upholstery)|custom furniture stores near me|custom upholstered|custom furniture shops$|"
     r"custom wood signs (?:wholesale|with preview)|custom (?:laser cut|light) signs|custom light signs|street sign maker|"
     r"custom wood plaques military|custom wood products near me|custom wood works near me|where can i get custom wood cuts|"
     r"custom dining room table pads|custom bench cushions|custom (?:granite|metal|glass)|custom dining table glass|"
     r"glass top|custom dining table chairs|custom decor dining table|custom furniture interior|custom fitted furniture",
     DIY, None, "Online or mass-produced custom products, or materials Grace does not work in (glass, metal, granite, laser-cut, lighted signs).")
# ---- 5. other places ------------------------------------------------------------------------------------------------
rule(r"(?:" + ELSEWHERE + r")", ELSE, None, "Names a place outside Grace's 15 East Texas towns.")
# ---- 6. East Texas: the town pages and hubs --------------------------------------------------------------------------
rule(r"furniture repair (?:tyler)|furniture restoration tyler", ET, "furniture-repair:tyler", "supporting")
rule(r"furniture repair (?:longview)", ET, "furniture-repair:longview", "supporting")
rule(r"furniture repair (?:henderson)", ET, "furniture-repair:henderson", "supporting")
rule(r"furniture repair (?:marshall)", ET, "furniture-repair:marshall", "supporting")
rule(r"furniture repair kilgore", ET, "furniture-repair:kilgore", "supporting")
rule(r"(?:custom )?cabinets? (?:in )?(?:tyler)|tyler cabinets|kitchen cabinets tyler|cabinets tyler", ET, "custom-cabinets:tyler", "supporting")
rule(r"cabinets? (?:in )?longview|cabinets longview", ET, "custom-cabinets:longview", "supporting")
rule(r"cabinets kilgore", ET, "custom-cabinets:kilgore", "supporting")
rule(r"custom furniture tyler", ET, "custom-furniture:tyler", "supporting")
rule(r"signs tyler", ET, "custom-signs:tyler", "supporting")
rule(r"woodworking longview", ET, "custom-furniture:longview", "supporting")
rule(r"cabinet makers (?:east )?texas|texas cabinet manufacturers|cabinets in texas|custom cabinets texas", SVC, "cc", "supporting")
rule(r"custom furniture texas", SVC, "cf", "supporting")
rule(r"custom signs texas", SVC, "cs", "supporting")
# ---- 7. Grace's services ---------------------------------------------------------------------------------------------
# cabinet painting: costs go to the guide, the service and its questions to the service page
rule(r"^cabinet painting prices$", SVC, "g-paintcost", "target")
rule(r"cabinet painting (?:price|prices|cost|estimate|quote)|how much (?:paint cabinets|is cabinet painting|does kitchen cabinet painting|"
     r"kitchen cabinet painting)|kitchen cabinets painting cost|is cabinet painting worth|cabinet painting cost per door|"
     r"cabinet painting cost calculator|how much does cabinet refinishing cost", Q, "g-paintcost", "supporting")
rule(r"^cabinet painting$", SVC, "paint", "target")
rule(r"cabinet painting|cabinet paintings|kitchen cabinets? (?:painting|repainting|refinishing)|cabinet painting kitchen|"
     r"cabinet paint (?:colors|black|white|grey|gray|green|blue|brown|espresso|finish|type|top coat|no prep)|"
     r"cabinet paint colors|can kitchen cabinets be (?:painted|spray painted|restained|stained)|is painting cabinets|"
     r"does painting cabinets|when painting cabinets|who paint kitchen cabinets|does cabinet paint need|"
     r"is cabinet paint oil|how paint cabinets white|kitchen cabinets refinishing near me",
     SVC, "paint", "supporting")
# cabinet refacing
rule(r"^what is cabinet resurfacing$", Q, "g-reface", "target")
rule(r"what (?:does )?(?:is )?cabinet refacing (?:mean|cost)|what is (?:cabinet resurfacing|refacing cabinet doors)|"
     r"how (?:does|is) (?:kitchen )?cabinet refacing|how to do cabinet refacing|how long does cabinet refacing|"
     r"is (?:cabinet )?refacing|refacing cabinets cheaper|cabinet refacing cost vs|what does cabinet refacing|"
     r"cabinet refacing calculator|kitchen cabinets resurfacing|what is cabinet refacing cost",
     Q, "g-reface", "supporting")
rule(r"^cabinet refacing$", SVC, "reface", "target")
rule(r"refacing|kitchen cabinets door replacement", SVC, "reface", "supporting")
# built-ins
rule(r"^built in shelves$", SVC, "builtin", "target")
rule(r"built ?in (?:shelves|shelving|bookshel|bookcase|record|shoe|garage|pantry|glass|metal)|bookshel|"
     r"how much (?:do )?built in|can built in cabinets", SVC, "builtin", "supporting")
# deck and porch (before the furniture cost rules, so deck cost questions stay on the deck page)
rule(r"^deck repair$", SVC, "deck", "target")
rule(r"deck", SVC, "deck", "supporting")
rule(r"^porch repair$", SVC, "porch", "target")
rule(r"porch", SVC, "porch", "supporting")
# chairs
rule(r"^chair caning$", SVC, "chair", "target")
rule(r"^how to repair chair caning$", Q, "g-cane", "target")
rule(r"how to repair chair (?:caning|seat)|what is chair caning|chair cane weaving|chair caning spline|"
     r"cane chair back replacement|how much does caning a chair|chair caning repair cost|cane chair repair cost",
     Q, "g-cane", "supporting")
rule(r"cane|caning|recaned|chair repair|chair leg|chair back|chair fixing|repair chair|chairs? .*repair|"
     r"how long should a chair last|can chair repair", SVC, "chair", "supporting")
# dining tables
rule(r"^custom dining table$", SVC, "dining", "target")
rule(r"^what size dining table seats 8$", Q, "g-table8", "target")
rule(r"dining table dimensions|space per person at dining|seats 8", Q, "g-table8", "supporting")
rule(r"dining (?:room )?table|custom (?:made )?dining|custom (?:hardwood|farmhouse|oval) dining|custom wood table|"
     r"custom height dining|custom (?:wood )?table (?:top|legs|base|design)|custom dining tables", SVC, "dining", "supporting")
# table refinishing
rule(r"^table refinishing$", SVC, "table", "target")
rule(r"refinishing (?:a table|antique table)|table (?:top )?refinishing|refinish table top|table (?:repair|restoration|resurfacing)|"
     r"wood table (?:finish )?repair|how long does it take to sand a table|painting table and chairs|table top refinishing|"
     r"veneer top", SVC, "table", "supporting")
# veneer
rule(r"^how to repair furniture veneer$", Q, "g-veneer", "target")
rule(r"veneer", Q, "g-veneer", "supporting")
# furniture refinishing and its costs
rule(r"^furniture refinishing cost$", SVC, "g-refcost", "target")
rule(r"(?:refinishing|repair|restoration) (?:cost|price guide)|how much (?:does )?(?:it cost to refinish|furniture (?:refinishing|repair|restoration))|"
     r"how much does furniture|wood furniture repair cost|how much does it cost to refinish|furniture repair (?:estimate|quote)|"
     r"furniture restoration quotes|chair repair (?:cost|price)", Q, "g-refcost", "supporting")
rule(r"^furniture refinishing$", SVC, "refinish", "target")
rule(r"refinish|refurbish|furniture refinishing|finish repair|furniture repair paint|custom furniture painting", SVC, "refinish", "supporting")
# antique
rule(r"^antique furniture restoration$", SVC, "antique", "target")
rule(r"antique|vintage furniture restoration|heirloom", SVC, "antique", "supporting")
# deck and porch
# signs
rule(r"^custom wood signs$", SVC, "cs", "target")
rule(r"\bsigns?\b|plaques", SVC, "cs", "supporting")
# custom cabinets
rule(r"^custom cabinets$", SVC, "cc", "target")
rule(r"custom cabinet|custom upper cabinets|semi custom cabinets|custom cabinetry|cabinet makers|cabinets? makers near me|"
     r"kitchen cabinets custom|custom cabinets kitchen|kitchen cabinets repair|why custom cabinets|"
     r"how much (?:are|should|is) custom cabinet|what (?:are|do) custom cabinets|what is custom cabinetry|"
     r"are custom cabinets|how long do custom cabinets|where to buy custom cabinets|kitchen cabinets near me$|"
     r"kitchen cabinets makers near me", SVC, "cc", "supporting")
# custom furniture
rule(r"^custom furniture$", SVC, "cf", "target")
rule(r"custom furniture (?:restoration|repair)|custom furniture refinishing", SVC, "fr", "supporting")
rule(r"custom (?:wood )?furniture|custom wood (?:end|side) table|why custom furniture|what is custom furniture|"
     r"where to get custom furniture|how much does custom furniture|custom furniture carpenter|furniture woodworking|"
     r"carpenter near me for furniture|handmade furniture", SVC, "cf", "supporting")
# furniture repair and restoration
rule(r"^furniture repair$", SVC, "fr", "target")
rule(r"furniture repair|furniture restoration|furniture restorer|restored furniture|wood furniture (?:repair|restoration|damage)|"
     r"who (?:fixes|repair|repairs) furniture|where to repair furniture|broken furniture repair|couch wood frame|"
     r"what is furniture restoration|how to repair furniture$|how to repair broken wood furniture",
     SVC, "fr", "supporting")
rule(r"when to replace furniture|how often (?:should you |to )?replace furniture|will furniture stores", Q, "fr", "supporting")
# home: the shop itself
rule(r"^woodworking shop$", SVC, "home", "target")
rule(r"woodworking shop|wood work shop|woodworking near|custom wood works|carpenter (?:work|for hire|local|near me|needed|handyman)|"
     r"^carpenter work$|^carpenter for hire$|^carpenter local$|how carpenter work|where carpenter work|wood shop near me",
     SVC, "home", "supporting")
rule(r"carpenter (?:hourly|rate)", SVC, None, "Asks for an hourly rate; Grace quotes each job and publishes no rates.")


LOCAL = re.compile(r"near me|nearby|in my area|\blocal\b|built ?in|bookshel")


def classify(kw, orig=""):
    # the sheet's "Other place" is right whenever the keyword is not a near-me search
    if orig == "Other place" and not LOCAL.search(kw) and not re.search(EAST_TX, kw):
        return ELSE, None, "Names a place outside Grace's 15 East Texas towns."
    for pat, fit, page, note in R:
        if pat.search(kw):
            return fit, page, note
    return None


def main():
    wb = openpyxl.load_workbook(XLSX, read_only=True)
    allr = [r for r in list(wb["ALL"].iter_rows(values_only=True))[1:] if (r[3] or 0) >= 10]
    fit_tab = {r[1] for r in list(wb["FIT FOR GRACE 10+"].iter_rows(values_only=True))[1:]}
    out, unmatched_fit, fallback = [], [], 0
    for seed, kw, typ, vol, sd, cpc, pd, orig in allr:
        kw = str(kw).strip()
        c = classify(kw.lower(), orig)
        if c is None:
            if kw in fit_tab:
                unmatched_fit.append(kw)
                continue
            # keywords the sheet already judged unfit and no rule claims keep the sheet's verdict
            fallback += 1
            c = (orig, None, "Sheet verdict kept: " + {
                "Other place": "names a place outside Grace's area.",
                "Off-topic": "not woodwork Grace does.",
                "DIY or product": "a product or do-it-yourself search.",
            }.get(orig, "not a fit."))
        fit, page, note = c
        if fit in (SVC, Q) and kw.lower().endswith("near me") or "near me" in kw.lower() and fit in (SVC, Q):
            fit = NEAR
        out.append(dict(seed=seed, kw=kw, type=typ, vol=vol, sd=sd, orig=orig, fit=fit, page=page,
                        role=note if page else "", reason="" if page else note))
    if unmatched_fit:
        print("Fit-tab keywords with no explicit rule:", len(unmatched_fit))
        for k in unmatched_fit:
            print("  ", k)
        return 1
    fits = (SVC, NEAR, Q, ET)
    # every target must be a real keyword on its page
    for key, kw in TARGETS.items():
        hit = [o for o in out if o["kw"].lower() == kw]
        if not hit or hit[0]["page"] != key or hit[0]["role"] != "target":
            print("target not assigned:", key, kw, hit[:1])
            return 1
    with open(ROOT / "docs" / "keyword-assignments.csv", "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["Keyword", "Volume", "SEO difficulty", "Sheet fit", "Corrected fit", "Page", "Role", "Unused reason"])
        for o in sorted(out, key=lambda o: (-o["vol"], o["kw"])):
            w.writerow([o["kw"], o["vol"], o["sd"], o["orig"], o["fit"], PAGES[o["page"]][0] if o["page"] else "",
                        o["role"], o["reason"]])
    write_md(out, fits, fallback)
    by = Counter(o["fit"] for o in out)
    print("rows", len(out), dict(by))
    print("assigned", sum(1 for o in out if o["page"]), "fit but unused",
          sum(1 for o in out if not o["page"] and o["fit"] in fits))
    return 0


def write_md(out, fits, fallback):
    changed = [o for o in out if (o["orig"] != o["fit"]) and not (o["orig"] == "Service (no place)" and o["fit"] == SVC)]
    pages = defaultdict(list)
    for o in out:
        if o["page"]:
            pages[o["page"]].append(o)
    md = ["# Keyword plan", "",
          "Source: keywords/Grace-Woodwork-Keywords.xlsx (Ubersuggest, United States, pulled Oct 9 2026), ALL tab, every",
          "keyword with 10 or more monthly searches. Built by `python3 scripts/keyword_plan.py`; the full row-by-row list",
          "is in docs/keyword-assignments.csv.", "",
          "## Read this first: national volume is not East Texas demand", "",
          "Ubersuggest gives United States totals. A keyword with no town in it (\"cabinet painting\", 40,500 a month) shows",
          "what people across the country type, not how many people in Kilgore, Longview or Tyler search for it. The East",
          "Texas share is a small fraction, and Google answers those searches with nearby businesses. The volumes rank",
          "the topics against each other; they are not a forecast of calls. The local keywords that actually name Grace's",
          "towns are tiny (\"furniture repair tyler tx\" 40, \"kitchen cabinets tyler tx\" 70, \"furniture repair longview tx\" 20),",
          "which is why the 60 town pages stay as they are and the new pages chase the service terms.", "",
          "## How column H was corrected", "",
          "Column H (\"Fit for Grace\") was a first sort. The script re-sorts every row with explicit rules. Main corrections:", "",
          "- \"near me\" searches were filed as Other place. They are local intent from whoever searches, so a near-me search",
          "  for a Grace service is now \"Service near me\" and assigned (for example \"furniture repair near me\" 18,100,",
          "  \"deck repair near me\" 9,900, \"cane chair repair near me\" 1,600).",
          "- \"built in shelves\" (4,400) and \"built in bookshelves\" (12,100) were filed as Other place; they are Grace's",
          "  built-in work and go to the built-in shelves page.",
          "- Several DIY or product rows are really service searches: \"kitchen cabinets repainting\" (27,100), \"cabinet",
          "  painting kitchen\" (22,200), \"custom cabinets kitchen\" (12,100), \"kitchen cabinets door replacement\" (12,100),",
          "  \"kitchen cabinets resurfacing\" (9,900), \"kitchen cabinets refinishing\" (6,600).",
          "- Many Service rows are not Grace at all: other companies (\"woodworking shop klingspor\", \"thomas johnson antique",
          "  furniture restoration\"), paint brands (\"cabinet paint from sherwin williams\"), hobby shop gear (\"woodworking shop",
          "  dust collection\"), upholstery and leather, and other cities (\"custom cabinets utah\"). These are now Other",
          "  business or brand, DIY or product, Off-topic or Other place.", "",
          f"Rows whose fit changed (other than Service to Service): {len(changed)}. Rows the sheet already called unfit and",
          f"no rule touched keep the sheet's verdict: {fallback}.", "",
          "## Summary", "",
          "| Corrected fit | Keywords | Assigned to a page | Deliberately unused |", "|---|---:|---:|---:|"]
    for fit in (SVC, NEAR, Q, ET, BRAND, DIY, ELSE, OFF):
        rows = [o for o in out if o["fit"] == fit]
        md.append(f"| {fit} | {len(rows)} | {sum(1 for o in rows if o['page'])} | {sum(1 for o in rows if not o['page'])} |")
    md += ["", "## Keywords by page", "",
           "Target = the main keyword the page is written for. Supporting = worked into headings, copy or FAQs where it",
           "reads naturally. Town pages keep their Sep 1 text; their keywords are recorded here, not added to the text.", ""]
    order = list(TARGETS) + sorted(k for k in pages if k not in TARGETS)
    for key in order:
        rows = sorted(pages.get(key, []), key=lambda o: (o["role"] != "target", -o["vol"]))
        if not rows:
            continue
        url, name = PAGES[key]
        md += [f"### {name}: `{url}`", "", "| Keyword | Volume | SD | Role |", "|---|---:|---:|---|"]
        md += [f"| {o['kw']} | {o['vol']:,} | {o['sd']} | {o['role']} |" for o in rows]
        md.append("")
    md += ["## Keywords that fit but are deliberately unused", "",
           "Fits Grace's trade but no page is written for it, with the reason.", "",
           "| Keyword | Volume | Corrected fit | Reason |", "|---|---:|---|---|"]
    for o in sorted((o for o in out if not o["page"] and o["fit"] in fits), key=lambda o: -o["vol"]):
        md.append(f"| {o['kw']} | {o['vol']:,} | {o['fit']} | {o['reason']} |")
    md += ["", "## Rows from the FIT FOR GRACE 10+ tab that turned out not to fit", "",
           "These were on Maurice's fit tab. Each one is listed with the reason it was moved out.", "",
           "| Keyword | Volume | Now | Reason |", "|---|---:|---|---|"]
    for o in sorted((o for o in out if not o["page"] and o["fit"] not in fits and o["orig"] in ("Service (no place)", "Question", "East Texas")),
                    key=lambda o: -o["vol"]):
        md.append(f"| {o['kw']} | {o['vol']:,} | {o['fit']} | {o['reason']} |")
    (ROOT / "docs" / "keyword-plan.md").write_text("\n".join(md) + "\n")


if __name__ == "__main__":
    sys.exit(main())
